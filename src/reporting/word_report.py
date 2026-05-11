from __future__ import annotations

# narratives şablonlarıyla DOCX rapor; tablolar ve görseller eklenir.

from pathlib import Path

from docx import Document
from docx.shared import Inches

from src.analysis.statistics import AnalysisResults
from src.reporting.narratives import build_report_narratives


def _format_table_cell(header: str, value: str, metric_value: str | None) -> str:
    raw = str(value)
    if raw in {"nan", "None"}:
        return raw
    header_l = header.lower()
    metric_l = (metric_value or "").lower()
    if header in {"Price"}:
        return f"{raw} TL"
    if header in {"Area"}:
        return f"{raw} m²"
    if header_l == "midpoint":
        return f"{raw} TL"
    if header_l in {"mean", "low", "high", "std", "margin"}:
        if metric_l == "price":
            return f"{raw} TL"
        if metric_l == "area":
            return f"{raw} m²"
    return raw


def _add_table(doc: Document, title: str, rows: list[list[str]]) -> None:
    doc.add_heading(title, level=2)
    if not rows:
        doc.add_paragraph("No data.")
        return
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for idx, value in enumerate(rows[0]):
        hdr_cells[idx].text = str(value)
    metric_idx = rows[0].index("metric") if "metric" in rows[0] else None
    for row in rows[1:]:
        cells = table.add_row().cells
        metric_value = str(row[metric_idx]) if metric_idx is not None and metric_idx < len(row) else None
        for idx, value in enumerate(row):
            header = str(rows[0][idx])
            cells[idx].text = _format_table_cell(header, str(value), metric_value)


def _add_part2_calculation_steps(doc: Document, ci_df) -> None:
    doc.add_heading("[Part 2 - Calculation Steps] — Statistical Notes", level=3)
    rows = [["metric", "confidence", "mean", "std", "n", "t_crit", "margin", "formula", "interval"]]
    for row in ci_df.itertuples(index=False):
        metric = str(getattr(row, "metric", "Metric"))
        confidence = float(getattr(row, "confidence", 0.95))
        mean = float(getattr(row, "mean", 0.0))
        std = float(getattr(row, "std", 0.0))
        n = int(round(float(getattr(row, "n", 0.0))))
        t_crit = float(getattr(row, "t_crit", 0.0))
        margin = float(getattr(row, "margin", 0.0))
        low = float(getattr(row, "low", 0.0))
        high = float(getattr(row, "high", 0.0))
        rows.append(
            [
                metric,
                f"{int(confidence * 100)}%",
                f"{mean:,.4f} {'TL' if metric == 'Price' else 'm²'}",
                f"{std:,.4f} {'TL' if metric == 'Price' else 'm²'}",
                str(n),
                f"{t_crit:,.4f}",
                f"{margin:,.4f} {'TL' if metric == 'Price' else 'm²'}",
                "mean ± t_crit * (std / sqrt(n))",
                f"[{low:,.4f}, {high:,.4f}] {'TL' if metric == 'Price' else 'm²'}",
            ]
        )
    _add_table(doc, "Part 2 - Confidence Interval Calculation Trace", rows)


def _add_interpretations(doc: Document, section_title: str, comments: list[str]) -> None:
    doc.add_heading(f"[{section_title}] — Statistical Interpretation", level=3)
    if not comments:
        paragraph = doc.add_paragraph("No interpretation available for this section.", style="Normal")
        paragraph.paragraph_format.space_after = 2
        return
    for comment in comments:
        chunks = [chunk.strip() for chunk in str(comment).split("\n\n") if chunk.strip()]
        if not chunks:
            chunks = [str(comment)]
        for chunk in chunks:
            paragraph = doc.add_paragraph(chunk, style="Normal")
            paragraph.paragraph_format.space_after = 3
    doc.add_paragraph("")


def _add_section_charts(doc: Document, title: str, charts: list[Path]) -> None:
    if not charts:
        return
    doc.add_heading(title, level=3)
    for chart_path in charts:
        if chart_path.exists():
            doc.add_paragraph(chart_path.name)
            doc.add_picture(str(chart_path), width=Inches(6))


def _charts_for_part(chart_paths: list[Path], part: str) -> list[Path]:
    selected: list[Path] = []
    for chart in chart_paths:
        name = chart.name.lower()
        if part == "part1" and ("price_hist" in name or "price_box" in name or "area_hist" in name):
            selected.append(chart)
        elif part == "part3" and name.startswith("part3_"):
            selected.append(chart)
        elif part == "part5" and "scatter_reg" in name:
            selected.append(chart)
        elif part == "part6" and ("anova" in name or "tukey" in name):
            selected.append(chart)
    return selected


def export_word_report(
    student_no: str,
    cities: tuple[str, str, str],
    results: AnalysisResults,
    chart_paths: list[Path],
    output_path: str | Path,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading("Statistics Term Project Report", 0)
    doc.add_paragraph(f"Student Number: {student_no}")
    doc.add_paragraph(f"Cities: {cities[0]}, {cities[1]}, {cities[2]}")
    doc.add_paragraph("Generated with Python statistical automation.")
    narratives = build_report_narratives(results, cities)

    _add_table(
        doc,
        "Part 1 - Descriptive Statistics",
        [results.descriptive.reset_index().columns.tolist()]
        + results.descriptive.reset_index().round(4).astype(str).values.tolist(),
    )
    _add_table(
        doc,
        "Part 1 - Price Frequency Distribution",
        [results.part1_price_frequency.columns.tolist()]
        + results.part1_price_frequency.round(4).astype(str).values.tolist(),
    )
    _add_interpretations(doc, "Part 1 - Descriptive Statistics", narratives.part1_descriptive)
    _add_section_charts(doc, "Part 1 - Charts", _charts_for_part(chart_paths, "part1"))
    _add_table(
        doc,
        "Part 2 - Confidence Intervals",
        [results.confidence_intervals.columns.tolist()]
        + results.confidence_intervals.round(4).astype(str).values.tolist(),
    )
    _add_part2_calculation_steps(doc, results.confidence_intervals)
    _add_interpretations(doc, "Part 2 - Confidence Intervals", narratives.part2_ci)
    _add_table(
        doc,
        "Part 3 - One Sample Tests",
        [results.one_sample_tests.columns.tolist()]
        + results.one_sample_tests.round(4).astype(str).values.tolist(),
    )
    _add_interpretations(doc, "Part 3 - One Sample Tests", narratives.part3_one_sample)
    _add_section_charts(doc, "Part 3 - Charts", _charts_for_part(chart_paths, "part3"))
    _add_table(
        doc,
        "Part 4 - Two Sample Tests",
        [results.two_sample_tests.columns.tolist()]
        + results.two_sample_tests.round(4).astype(str).values.tolist(),
    )
    _add_interpretations(doc, "Part 4 - Two Sample Tests", narratives.part4_two_sample)
    _add_table(
        doc,
        "Part 5 - Regression Summary",
        [results.regression_summary.columns.tolist()]
        + results.regression_summary.round(4).astype(str).values.tolist(),
    )
    _add_interpretations(doc, "Part 5 - Regression", narratives.part5_regression)
    _add_section_charts(doc, "Part 5 - Charts", _charts_for_part(chart_paths, "part5"))
    _add_table(
        doc,
        "Part 6 - ANOVA",
        [results.anova_table.columns.tolist()] + results.anova_table.round(4).astype(str).values.tolist(),
    )
    _add_interpretations(doc, "Part 6 - ANOVA", narratives.part6_anova)
    _add_table(
        doc,
        "Part 6 - Tukey Multiple Comparison",
        [results.tukey_table.columns.tolist()] + results.tukey_table.astype(str).values.tolist(),
    )
    _add_interpretations(doc, "Part 6 - Tukey Multiple Comparison", narratives.part6_tukey)
    _add_section_charts(doc, "Part 6 - Charts", _charts_for_part(chart_paths, "part6"))

    _add_interpretations(doc, "Overall Evaluation", narratives.general_summary)

    doc.save(output_path)
    return output_path
